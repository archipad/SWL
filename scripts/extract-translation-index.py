"""Convertit l'index EN -> FR fourni par la communauté en dictionnaire web.

Usage:
  python scripts/extract-translation-index.py "Index de traduction ...xlsx"
"""

from __future__ import annotations

import json
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl


def text(value: object) -> str:
    return str(value or "").strip()


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Indiquez le chemin du fichier XLSX de traduction.")
    source = Path(sys.argv[1]).resolve()
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["EN -> FR"]
    updated = sheet["B3"].value
    if isinstance(updated, (date, datetime)):
        updated = updated.date().isoformat() if isinstance(updated, datetime) else updated.isoformat()

    game_terms: dict[str, str] = {}
    keywords: dict[str, dict[str, str]] = {}
    for row in range(13, sheet.max_row + 1):
        en_term, fr_term = text(sheet.cell(row, 1).value), text(sheet.cell(row, 2).value)
        if en_term and fr_term:
            game_terms[en_term] = fr_term
        en_keyword, fr_keyword = text(sheet.cell(row, 6).value), text(sheet.cell(row, 7).value)
        if en_keyword and fr_keyword:
            keywords[en_keyword] = {"fr": fr_keyword, "type": text(sheet.cell(row, 8).value)}

    payload = {
        "schemaVersion": 1,
        "source": source.name,
        "updatedAt": updated,
        "gameTerms": dict(sorted(game_terms.items())),
        "keywords": dict(sorted(keywords.items())),
    }
    root = Path(__file__).resolve().parents[1]
    json_path = root / "src" / "data" / "translationGlossary.json"
    js_path = root / "public" / "assistant" / "translation-glossary.js"
    serialized = json.dumps(payload, ensure_ascii=False, indent=2) + "\n"
    json_path.write_text(serialized, encoding="utf-8")
    js_path.write_text("window.SWL_TRANSLATION_GLOSSARY=" + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
    print(f"{len(game_terms)} termes et {len(keywords)} mots-clés exportés.")


if __name__ == "__main__":
    main()
